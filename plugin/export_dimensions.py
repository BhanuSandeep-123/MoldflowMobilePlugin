"""
export_dimensions.py
--------------------
Extract a Moldflow part's DESIGN DIMENSIONS from the active study and write them
to a formatted Excel workbook:

  * Bounding box  (X/Y/Z min/max/size, overall size, diagonal) + mesh volume
  * Wall thickness stats (min / avg / max) via DiagnosisManager.GetThicknessDiagnosis
  * Cylindrical through-hole diameters & positions (mesh_geometry.detect_holes)

Fast path: one bulk Synergy.Project().ExportModel() to a .udm text file, parsed
in pure Python -- no per-node COM calls.

REQUIRES A MESHED STUDY (a raw imported STEP/CAD solid has no nodes).

Run from Synergy via run_plugin.vbs (Macro 2), or standalone:
    python export_dimensions.py
"""

from __future__ import annotations

import datetime
import os
import sys
import tempfile
import traceback

import mesh_geometry as mg

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "design_dimensions.xlsx")
LOG = os.path.join(HERE, "dimensions_log.txt")


def _length_unit(units: str) -> str:
    return "in" if str(units).lower().startswith("eng") else "mm"


def _m_to_display_factor(units: str) -> float:
    # ExportModel writes coordinates in SI base units (METERS). Convert to the
    # display length unit (mm for Metric, in for English).
    return 1000.0 / 25.4 if str(units).lower().startswith("eng") else 1000.0


def mesh_summary(syn):
    try:
        return syn.DiagnosisManager().GetMeshSummary(False)
    except Exception:
        return None


def thickness_stats(syn, bbox_thickness_disp, units):
    """min/avg/max wall thickness in display units, or None.

    The diagnosis value's unit (SI meters vs display) is auto-detected by
    checking which interpretation lands closest to the known bounding-box
    thickness (smallest bbox dimension, already in display units).
    """
    try:


        
        diag = syn.DiagnosisManager()
        ints = syn.CreateIntegerArray()
        dbls = syn.CreateDoubleArray()
        count = diag.GetThicknessDiagnosis(0.0, 0.0, ints, dbls)  # 0,0 => all
        if not count:
            return None
        vals = [v for v in dbls.ToVBSArray() if v == v and v > 0]  # drop NaN/0
    except Exception:
        return None
    if not vals:
        return None

    raw_max = max(vals)
    f_disp = _m_to_display_factor(units)
    if abs(raw_max * f_disp - bbox_thickness_disp) < abs(raw_max - bbox_thickness_disp):
        f = f_disp          # values were in meters
    else:
        f = 1.0             # values already in display units
    vals = [v * f for v in vals]
    return {
        "count": len(vals),
        "min": min(vals),
        "avg": sum(vals) / len(vals),
        "max": max(vals),
    }


def analyze(syn, expected_nodes):
    """
    Export the mesh once, then compute bbox + holes (in display units).
    Returns (result_dict, warning) or (None, None) if no mesh.
    """
    if not expected_nodes:
        return None, None

    fd, udm = tempfile.mkstemp(suffix=".udm", prefix="moldflow_dims_")
    os.close(fd)
    try:
        syn.Project().ExportModel(udm)
        coords, tets, tris, bbox_m = mg.parse_udm(udm)   # coords in METERS
        # Hole detection needs the OUTER SURFACE. For a 3D solid mesh that means
        # extracting the boundary of the tetrahedra; for fusion all nodes qualify.
        surf = mg.surface_coords(coords, tets, tris)
        holes_m = mg.detect_holes(surf, bbox_m)
    finally:
        try:
            os.remove(udm)
        except OSError:
            pass

    if not coords:
        return None, None

    units = syn.GetUnits()
    f = _m_to_display_factor(units)

    bbox = {a: (bbox_m[a][0] * f, bbox_m[a][1] * f,
                (bbox_m[a][1] - bbox_m[a][0]) * f) for a in mg.AXES}

    pa, pb = holes_m["planar"]
    holes = []
    for h in holes_m["holes"]:
        holes.append({
            "diameter": h["diameter"] * f,
            "center": {pa: h["center"][pa] * f, pb: h["center"][pb] * f},
            "npts": h["npts"],
        })

    warning = None
    if expected_nodes and len(coords) != expected_nodes:
        warning = (f"Parsed {len(coords)} nodes but mesh summary reported "
                   f"{expected_nodes}.")

    result = {
        "nodes": len(coords),
        "is_3d": bool(tets),
        "surface_nodes": len(surf),
        "bbox": bbox,
        "hole_axis": holes_m["axis"],
        "hole_plane": (pa, pb),
        "holes": holes,
    }
    return result, warning


def write_excel(study_name, units, res, thickness, volume, thickness_note=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    ulen = _length_unit(units)
    bbox = res["bbox"]
    pa, pb = res["hole_plane"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Design Dimensions"

    title_font = Font(size=14, bold=True)
    sect_font = Font(size=12, bold=True, color="2F5496")
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2F5496")
    label_font = Font(bold=True)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    right = Alignment(horizontal="right")
    left = Alignment(horizontal="left")

    r = 1
    ws.cell(r, 1, "Design Dimensions Report").font = title_font
    r += 2

    for k, v in [
        ("Study", study_name),
        ("Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Unit system", f"{units} ({ulen})"),
        ("Mesh nodes", res["nodes"]),
    ]:
        ws.cell(r, 1, k).font = label_font
        ws.cell(r, 2, v)
        r += 1

    def section(title):
        nonlocal r
        r += 1
        ws.cell(r, 1, title).font = sect_font
        r += 1

    def table_header(cols):
        nonlocal r
        for c, h in enumerate(cols, start=1):
            cell = ws.cell(r, c, h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.border = border
            cell.alignment = left if c == 1 else right
        r += 1

    # ---- Bounding box ----
    section("Bounding Box")
    table_header(["Axis", f"Min ({ulen})", f"Max ({ulen})", f"Size ({ulen})"])
    for axis in ("x", "y", "z"):
        lo, hi, size = bbox[axis]
        ws.cell(r, 1, axis.upper()).font = label_font
        ws.cell(r, 1).border = border
        for c, val in enumerate((lo, hi, size), start=2):
            cell = ws.cell(r, c, round(val, 3))
            cell.number_format = "0.000"
            cell.border = border
            cell.alignment = right
        r += 1
    r += 1
    xs, ys, zs = bbox["x"][2], bbox["y"][2], bbox["z"][2]
    diag = (xs ** 2 + ys ** 2 + zs ** 2) ** 0.5
    rows = [
        (f"Overall size (X x Y x Z) [{ulen}]", f"{xs:.3f} x {ys:.3f} x {zs:.3f}"),
        (f"Bounding-box diagonal [{ulen}]", round(diag, 3)),
    ]
    if volume is not None:
        rows.append(("Mesh volume (cm^3)", round(volume, 3)))
    for k, v in rows:
        ws.cell(r, 1, k).font = label_font
        ws.cell(r, 2, v)
        r += 1

    # ---- Wall thickness ----
    section("Wall Thickness")
    # Always show the nominal thickness (smallest bounding dimension).
    nominal = min(bbox[a][2] for a in ("x", "y", "z"))
    ws.cell(r, 1, f"Nominal thickness (min bounding dim) [{ulen}]").font = label_font
    ws.cell(r, 2, round(nominal, 3))
    r += 1
    if thickness_note:
        ws.cell(r, 1, thickness_note)
        r += 1
    if thickness is None:
        if not thickness_note:
            ws.cell(r, 1, "No thickness diagnosis available "
                          "(needs a fusion/midplane mesh).")
            r += 1
    else:
        table_header(["Statistic", f"Value ({ulen})"])
        for k, v in [("Minimum", thickness["min"]),
                     ("Average", thickness["avg"]),
                     ("Maximum", thickness["max"])]:
            ws.cell(r, 1, k).font = label_font
            ws.cell(r, 1).border = border
            cell = ws.cell(r, 2, round(v, 4))
            cell.number_format = "0.000"
            cell.border = border
            cell.alignment = right
            r += 1
        ws.cell(r, 1, "Elements measured").font = label_font
        ws.cell(r, 2, thickness["count"])
        r += 1

    # ---- Holes ----
    section(f"Holes (through {res['hole_axis'].upper()} axis)")
    holes = res["holes"]
    if not holes:
        ws.cell(r, 1, "No cylindrical through-holes detected.")
        r += 1
    else:
        table_header(["Hole #", f"Diameter ({ulen})",
                      f"Center {pa.upper()} ({ulen})",
                      f"Center {pb.upper()} ({ulen})"])
        for i, h in enumerate(holes, start=1):
            ws.cell(r, 1, i).font = label_font
            ws.cell(r, 1).border = border
            vals = (h["diameter"], h["center"][pa], h["center"][pb])
            for c, val in enumerate(vals, start=2):
                cell = ws.cell(r, c, round(val, 3))
                cell.number_format = "0.000"
                cell.border = border
                cell.alignment = right
            r += 1
        ws.cell(r, 1, "Holes found").font = label_font
        ws.cell(r, 2, len(holes))
        r += 1

    ws.column_dimensions["A"].width = 32
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 18

    wb.save(XLSX)


def main() -> int:
    from synergy_connect import get_synergy
    try:
        syn = get_synergy(allow_launch=False)
        sd = syn.StudyDoc()
        if sd is None:
            print("No active study. Open/activate your part study and retry.")
            return 2

        study_name = sd.StudyName()
        units = syn.GetUnits()
        print(f"Active study : {study_name}")
        print(f"Units        : {units}")

        summary = mesh_summary(syn)
        expected_nodes = int(summary.NodesCount()) if summary is not None else 0
        volume = summary.MeshVolume() if summary is not None else None

        if not expected_nodes:
            msg = (
                f"Study '{study_name}' has NO MESH (0 nodes). A freshly imported\n"
                "STEP/CAD solid must be meshed first: Mesh tab -> Generate Mesh."
            )
            print(msg)
            with open(LOG, "w", encoding="utf-8") as fh:
                fh.write(msg + "\n")
            return 3

        print(f"Meshed nodes : {expected_nodes} (exporting + analyzing...)")
        res, warning = analyze(syn, expected_nodes)
        if res is None:
            print("Mesh export produced no parseable nodes.")
            return 4
        if warning:
            print("WARNING:", warning)

        ulen = _length_unit(units)
        for axis in ("x", "y", "z"):
            lo, hi, size = res["bbox"][axis]
            print(f"  {axis.upper()}: {lo:.3f} .. {hi:.3f}  size={size:.3f} {ulen}")

        bbox_thickness = min(res["bbox"][a][2] for a in ("x", "y", "z"))
        thickness = None
        thickness_note = None
        if res["is_3d"]:
            # Thickness Diagnosis is a fusion/midplane concept; on a 3D solid
            # mesh it returns meaningless values (e.g. ray lengths). Skip it.
            thickness_note = ("3D solid mesh: per-element wall-thickness diagnosis "
                              "is not meaningful. Use a Dual Domain (fusion) mesh "
                              "for true wall-thickness stats.")
            print(f"  Thickness: skipped (3D mesh); nominal = {bbox_thickness:.3f} {ulen}")
        else:
            thickness = thickness_stats(syn, bbox_thickness, units)
            if thickness:
                print(f"  Thickness: min={thickness['min']:.3f} "
                      f"avg={thickness['avg']:.3f} max={thickness['max']:.3f} {ulen}")
            else:
                print("  Thickness: (no diagnosis available)")

        print(f"  Holes found: {len(res['holes'])}")
        for i, h in enumerate(res["holes"], start=1):
            pa, pb = res["hole_plane"]
            print(f"    #{i}: dia={h['diameter']:.3f} {ulen}  "
                  f"center {pa.upper()}={h['center'][pa]:.3f} "
                  f"{pb.upper()}={h['center'][pb]:.3f}")

        write_excel(study_name, units, res, thickness, volume, thickness_note)
        print(f"\nWrote {XLSX}")
        try:
            os.startfile(XLSX)
        except Exception:
            pass
        return 0

    except Exception:
        tb = traceback.format_exc()
        print("FAILED:\n" + tb)
        with open(LOG, "w", encoding="utf-8") as fh:
            fh.write("FAILED " + datetime.datetime.now().isoformat() + "\n\n" + tb)
        return 1


if __name__ == "__main__":
    sys.exit(main())
